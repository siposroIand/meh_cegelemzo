from typing import List

import pandas as pd

from core.utils import format_phone, parse_hu_numeric, split_address_parts


def build_detail_export(master: pd.DataFrame, source_df: pd.DataFrame) -> pd.DataFrame:
    src_cols = [
        "CEGNEV", "KUJ", "SZEKHELY", "ADOSZAM",
        "SRC_EMAIL", "SRC_TELEFON", "SRC_WEB",
        "SRC_SAJAT_TOKE", "SRC_NETTO_ARBEV"
    ]
    src = source_df.copy()
    for c in src_cols:
        if c not in src.columns:
            src[c] = ""
    src = src[src_cols].drop_duplicates(subset=["CEGNEV", "KUJ"])

    out = master.merge(
        src,
        how="left",
        left_on=["CEGNEV", "SOURCE_KUJ"],
        right_on=["CEGNEV", "KUJ"],
    )

    parts = out["SZEKHELY"].apply(split_address_parts).apply(pd.Series)

    out["SRC_TELEFON"] = out["SRC_TELEFON"].apply(format_phone)
    out["Mennyiség (kg)"] = pd.to_numeric(out["OSSZES_MENNYISEG"], errors="coerce").fillna(0).round(2)
    out["Mennyiség (t)"] = (pd.to_numeric(out["OSSZES_MENNYISEG"], errors="coerce").fillna(0) / 1000).round(4)
    out["Veszélyes"] = out["HULLADEK_VESZELYES_10"].apply(
        lambda x: "Igen" if pd.to_numeric(x, errors="coerce") == 1 else "Nem"
    )

    out["SRC_SAJAT_TOKE_NUM"] = parse_hu_numeric(out["SRC_SAJAT_TOKE"]).fillna(0).round(0).astype("Int64")
    out["SRC_NETTO_ARBEV_NUM"] = parse_hu_numeric(out["SRC_NETTO_ARBEV"]).fillna(0).round(0).astype("Int64")

    out = pd.concat([out, parts], axis=1)

    out = out[[
        "CEGNEV", "Vármegye", "IRSZ", "Település", "Cím",
        "SRC_EMAIL", "SRC_TELEFON",
        "ADOSZAM", "SOURCE_KUJ", "SRC_SAJAT_TOKE_NUM", "SRC_NETTO_ARBEV_NUM",
        "EV", "ADATTIPUS_LABEL", "HULLADEKKOD", "HULLADEK_MEGNEVEZES", "PRODUCT_GROUP",
        "Veszélyes", "Mennyiség (kg)", "Mennyiség (t)"
    ]].copy()

    out.rename(columns={
        "CEGNEV": "Cégnév",
        "SRC_EMAIL": "Email",
        "SRC_TELEFON": "Telefonszám",
        "ADOSZAM": "Adószám",
        "SOURCE_KUJ": "KÜJ",
        "SRC_SAJAT_TOKE_NUM": "Saját tőke",
        "SRC_NETTO_ARBEV_NUM": "Nettó árbevétel",
        "EV": "Év",
        "ADATTIPUS_LABEL": "Adattípus",
        "HULLADEKKOD": "Hulladékkód",
        "HULLADEK_MEGNEVEZES": "Hulladék megnevezés",
        "PRODUCT_GROUP": "Termékcsoport",
    }, inplace=True)

    return out.sort_values(
        ["Cégnév", "Év", "Hulladékkód", "Termékcsoport"],
        ascending=[True, False, True, True]
    ).reset_index(drop=True)


def build_detail_pivot_export(master: pd.DataFrame) -> pd.DataFrame:
    if master is None or master.empty:
        return pd.DataFrame(columns=["Cég", "Év", "Termékcsoport", "HAK", "Mennyiség (t)"])

    df = master.copy()
    df["Mennyiség (t)"] = pd.to_numeric(df["OSSZES_MENNYISEG"], errors="coerce").fillna(0) / 1000

    grouped = (
        df.groupby(["CEGNEV", "EV", "PRODUCT_GROUP", "HULLADEKKOD"], as_index=False)["Mennyiség (t)"]
        .sum()
        .sort_values(["CEGNEV", "EV", "PRODUCT_GROUP", "HULLADEKKOD"])
        .reset_index(drop=True)
    )

    grouped.rename(columns={
        "CEGNEV": "Cég",
        "EV": "Év",
        "PRODUCT_GROUP": "Termékcsoport",
        "HULLADEKKOD": "HAK",
    }, inplace=True)

    return grouped


def build_summary_export(master: pd.DataFrame, source_df: pd.DataFrame, groups: List[str]) -> pd.DataFrame:
    if master is None or master.empty:
        return pd.DataFrame()

    src_cols = [
        "CEGNEV", "KUJ", "SZEKHELY", "ADOSZAM",
        "SRC_EMAIL", "SRC_TELEFON", "SRC_SAJAT_TOKE", "SRC_NETTO_ARBEV"
    ]
    src = source_df.copy()
    for c in src_cols:
        if c not in src.columns:
            src[c] = ""
    src = src[src_cols].drop_duplicates(subset=["CEGNEV", "KUJ"])

    df = master.copy()
    df["Mennyiség (t)"] = pd.to_numeric(df["OSSZES_MENNYISEG"], errors="coerce").fillna(0) / 1000

    pivot = (
        df.groupby(["CEGNEV", "SOURCE_KUJ", "EV", "PRODUCT_GROUP"], as_index=False)["Mennyiség (t)"]
        .sum()
        .pivot_table(
            index=["CEGNEV", "SOURCE_KUJ", "EV"],
            columns="PRODUCT_GROUP",
            values="Mennyiség (t)",
            aggfunc="sum",
            fill_value=0
        )
        .reset_index()
    )

    total_base = (
        df[["CEGNEV", "SOURCE_KUJ", "EV", "HULLADEKKOD", "ADATTIPUS_LABEL", "OSSZES_MENNYISEG"]]
        .drop_duplicates()
        .copy()
    )
    total_base["Mennyiség (t)"] = pd.to_numeric(total_base["OSSZES_MENNYISEG"], errors="coerce").fillna(0) / 1000

    total_by_year = (
        total_base.groupby(["CEGNEV", "SOURCE_KUJ", "EV"], as_index=False)["Mennyiség (t)"]
        .sum()
        .rename(columns={"Mennyiség (t)": "Össz hulladék (tonna)"})
    )

    out = total_by_year.merge(pivot, on=["CEGNEV", "SOURCE_KUJ", "EV"], how="left")
    out = out.merge(
        src,
        how="left",
        left_on=["CEGNEV", "SOURCE_KUJ"],
        right_on=["CEGNEV", "KUJ"]
    )

    parts = out["SZEKHELY"].apply(split_address_parts).apply(pd.Series)
    out = pd.concat([out, parts], axis=1)

    out["Telefonszám"] = out["SRC_TELEFON"].apply(format_phone)
    out["SRC_SAJAT_TOKE_NUM"] = parse_hu_numeric(out["SRC_SAJAT_TOKE"]).fillna(0).round(0).astype("Int64")
    out["SRC_NETTO_ARBEV_NUM"] = parse_hu_numeric(out["SRC_NETTO_ARBEV"]).fillna(0).round(0).astype("Int64")

    for g in groups:
        if g not in out.columns:
            out[g] = 0.0

    final_cols = [
        "CEGNEV",
        "Vármegye",
        "IRSZ",
        "Település",
        "Cím",
        "SRC_EMAIL",
        "Telefonszám",
        "ADOSZAM",
        "SOURCE_KUJ",
        "SRC_SAJAT_TOKE_NUM",
        "SRC_NETTO_ARBEV_NUM",
        "EV",
        "Össz hulladék (tonna)",
    ] + groups

    out = out[final_cols].copy()
    out.rename(columns={
        "CEGNEV": "Cégnév",
        "SRC_EMAIL": "Email",
        "ADOSZAM": "Adószám",
        "SOURCE_KUJ": "KÜJ",
        "SRC_SAJAT_TOKE_NUM": "Saját tőke",
        "SRC_NETTO_ARBEV_NUM": "Nettó árbevétel",
        "EV": "Év",
    }, inplace=True)

    numeric_cols = ["Össz hulladék (tonna)"] + groups
    for c in numeric_cols:
        if c in out.columns:
            out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0.0)

    return out.sort_values(["Cégnév", "Év"], ascending=[True, False]).reset_index(drop=True)


def style_excel_worksheet(ws):
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

    header_fill = PatternFill(start_color="404040", end_color="404040", fill_type="solid")
    sub_fill_odd = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    sub_fill_even = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="BFBFBF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill = sub_fill_odd if row_idx % 2 == 0 else sub_fill_even
        for cell in row:
            cell.fill = fill
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for col in ws.columns:
        letter = col[0].column_letter
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 38)


def build_outline_report_sheet(ws, grouped_df: pd.DataFrame):
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

    thin = Side(style="thin", color="BFBFBF")
    bold_font = Font(bold=True)
    fill_company = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    fill_year = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    fill_group = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid")

    ws.sheet_properties.outlinePr.summaryBelow = True

    headers = ["Szint", "Megnevezés", "Év", "HAK", "Mennyiség (t)"]
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="404040", end_color="404040", fill_type="solid")
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        c.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 2

    if grouped_df is None or grouped_df.empty:
        return

    for company, df_company in grouped_df.groupby("Cég", sort=True):
        company_total = df_company["Mennyiség (t)"].sum()

        ws.cell(row=row_idx, column=1, value="Cég")
        ws.cell(row=row_idx, column=2, value=company)
        ws.cell(row=row_idx, column=5, value=float(company_total))
        for c in range(1, 6):
            ws.cell(row=row_idx, column=c).font = bold_font
            ws.cell(row=row_idx, column=c).fill = fill_company
            ws.cell(row=row_idx, column=c).border = Border(left=thin, right=thin, top=thin, bottom=thin)
        company_row = row_idx
        row_idx += 1

        company_start = row_idx

        for year, df_year in df_company.groupby("Év", sort=True):
            year_total = df_year["Mennyiség (t)"].sum()

            ws.cell(row=row_idx, column=1, value="Év")
            ws.cell(row=row_idx, column=2, value=company)
            ws.cell(row=row_idx, column=3, value=int(year) if pd.notna(year) else "")
            ws.cell(row=row_idx, column=5, value=float(year_total))
            for c in range(1, 6):
                ws.cell(row=row_idx, column=c).fill = fill_year
                ws.cell(row=row_idx, column=c).border = Border(left=thin, right=thin, top=thin, bottom=thin)
            year_row = row_idx
            row_idx += 1

            year_start = row_idx

            for group, df_group in df_year.groupby("Termékcsoport", sort=True):
                group_total = df_group["Mennyiség (t)"].sum()

                ws.cell(row=row_idx, column=1, value="Termékcsoport")
                ws.cell(row=row_idx, column=2, value=group)
                ws.cell(row=row_idx, column=3, value=int(year) if pd.notna(year) else "")
                ws.cell(row=row_idx, column=5, value=float(group_total))
                for c in range(1, 6):
                    ws.cell(row=row_idx, column=c).fill = fill_group
                    ws.cell(row=row_idx, column=c).border = Border(left=thin, right=thin, top=thin, bottom=thin)
                group_row = row_idx
                row_idx += 1

                group_start = row_idx

                for _, r in df_group.iterrows():
                    ws.cell(row=row_idx, column=1, value="HAK")
                    ws.cell(row=row_idx, column=2, value=r["Termékcsoport"])
                    ws.cell(row=row_idx, column=3, value=int(r["Év"]) if pd.notna(r["Év"]) else "")
                    ws.cell(row=row_idx, column=4, value=r["HAK"])
                    ws.cell(row=row_idx, column=5, value=float(r["Mennyiség (t)"]))
                    for c in range(1, 6):
                        ws.cell(row=row_idx, column=c).border = Border(left=thin, right=thin, top=thin, bottom=thin)
                    ws.row_dimensions[row_idx].outlineLevel = 3
                    row_idx += 1

                if group_start <= row_idx - 1:
                    for rr in range(group_start, row_idx):
                        ws.row_dimensions[rr].hidden = True
                    ws.row_dimensions[group_row].outlineLevel = 2

            if year_start <= row_idx - 1:
                for rr in range(year_start, row_idx):
                    ws.row_dimensions[rr].hidden = True
                ws.row_dimensions[year_row].outlineLevel = 1

        if company_start <= row_idx - 1:
            for rr in range(company_start, row_idx):
                ws.row_dimensions[rr].hidden = True
            ws.row_dimensions[company_row].outlineLevel = 0

    for row in ws.iter_rows(min_row=2):
        row[4].number_format = '#,##0.0000'

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 18